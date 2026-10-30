from __future__ import annotations

import html
import json
import shutil
import tempfile
from dataclasses import asdict, replace
from pathlib import Path

from PySide6.QtCore import QSignalBlocker, Qt, QUrl
from PySide6.QtGui import QAction, QDesktopServices, QIcon, QKeySequence, QPixmap
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFileDialog,
    QFormLayout, QHBoxLayout, QHeaderView, QInputDialog, QLabel, QLineEdit,
    QMainWindow, QMessageBox, QPlainTextEdit, QProgressBar, QPushButton,
    QScrollArea, QSplitter, QTableWidget, QTableWidgetItem, QTabWidget,
    QTextBrowser, QVBoxLayout, QWidget,
)

from .. import __version__
from ..core.dcr import assign_column, parse_column, read_range
from ..core.models import ChannelMapping, DcrRow, ReportConfig, ReportSettings
from ..core.project import load_project, save_project, validate_config
from ..jobs import JobRunner
from .theme import apply_theme, section_title


def text_item(value, editable=True):
    item = QTableWidgetItem(str(value) if value is not None else "")
    if not editable:
        item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
    return item


def check_item(value):
    item = QTableWidgetItem()
    item.setFlags(Qt.ItemFlag.ItemIsEnabled | Qt.ItemFlag.ItemIsSelectable | Qt.ItemFlag.ItemIsUserCheckable)
    item.setCheckState(Qt.CheckState.Checked if value else Qt.CheckState.Unchecked)
    return item


def make_table(headers):
    table = QTableWidget(0, len(headers))
    table.setHorizontalHeaderLabels(headers)
    table.setAlternatingRowColors(True)
    table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
    table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
    table.horizontalHeader().setStretchLastSection(True)
    table.verticalHeader().setVisible(False)
    return table


def button(label, callback, primary=False):
    control = QPushButton(label)
    control.setProperty("primary", primary)
    control.clicked.connect(callback)
    return control


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"SPD Report Generator {__version__}")
        self.setWindowIcon(QIcon(str(Path(__file__).resolve().parents[1] / "assets" / "app.ico")))
        available = QApplication.primaryScreen().availableGeometry()
        self.resize(min(1280, available.width()-40), min(850, available.height()-60))
        self.setMinimumSize(720, 460)
        self.config = ReportConfig()
        self.spd = self.ts = None
        self.project_path = None
        self.loaded_paths = {"spd": "", "touchstone": ""}
        self.pending_config = None
        self.loading_table = False
        self.action = ""
        self.preview_images = []
        self.temp = tempfile.TemporaryDirectory(prefix="spd-report-")
        self.preview_book = None
        self.runner = JobRunner(self)
        self.runner.progress.connect(self.on_progress)
        self.runner.finished.connect(self.on_finished)
        self.runner.failed.connect(self.on_failed)
        self.runner.busy_changed.connect(self.on_busy)
        self.make_menu()
        self.tabs = QTabWidget()
        input_scroll = QScrollArea()
        input_scroll.setWidgetResizable(True)
        input_scroll.setWidget(self.input_page())
        self.tabs.addTab(input_scroll, "1  입력")
        self.tabs.addTab(self.mapping_page(), "2  채널·그룹")
        self.tabs.addTab(self.cap_dcr_page(), "3  CAP·DCR")
        self.tabs.addTab(self.preview_page(), "4  미리보기·출력")
        root = QWidget()
        layout = QVBoxLayout(root)
        layout.addWidget(self.tabs, 1)
        status_row = QHBoxLayout()
        self.status = QLabel("SPD·Touchstone 파일을 선택하세요. DCR만 입력할 때는 채널을 추가하세요.")
        self.status.setWordWrap(True)
        self.progress = QProgressBar()
        self.progress.setMaximumWidth(210)
        self.progress.setRange(0, 100)
        self.cancel_button = button("취소", self.runner.cancel)
        self.cancel_button.setEnabled(False)
        status_row.addWidget(self.status, 1)
        status_row.addWidget(self.progress)
        status_row.addWidget(self.cancel_button)
        layout.addLayout(status_row)
        self.log = QPlainTextEdit()
        self.log.setObjectName("output_log")
        self.log.setReadOnly(True)
        self.log.setMaximumHeight(85)
        self.log.setPlaceholderText("입력 상태와 제외된 페이지의 사유가 여기에 표시됩니다.")
        layout.addWidget(self.log)
        self.setCentralWidget(root)
        apply_theme(self)
        self.saved_state = self.state_text()

    def make_menu(self):
        menu = self.menuBar().addMenu("파일")
        for label, shortcut, callback in (("새 프로젝트", "Ctrl+N", self.new_project), ("프로젝트 열기…", "Ctrl+O", self.open_project), ("프로젝트 저장…", "Ctrl+S", self.save_project)):
            action = QAction(label, self)
            action.setShortcut(QKeySequence(shortcut))
            action.triggered.connect(callback)
            menu.addAction(action)
        help_menu = self.menuBar().addMenu("Help")
        help_action = QAction("사용 설명서", self)
        help_action.setShortcut(QKeySequence("F1"))
        help_action.triggered.connect(self.open_help)
        help_menu.addAction(help_action)
        about = QAction("프로그램 정보", self)
        about.triggered.connect(lambda: QMessageBox.information(self, "프로그램 정보", f"SPD Report Generator {__version__}\nMIT License\n로컬 파일만 처리합니다.\nhttps://github.com/yunhyok/report-generator"))
        help_menu.addAction(about)

    def input_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.addWidget(section_title("입력 파일 · 각 파일은 선택 사항입니다"))
        self.spd_path = QLineEdit()
        self.ts_path = QLineEdit()
        for label, field, filters in (("SPD", self.spd_path, "SPD (*.spd);;모든 파일 (*)"), ("Touchstone", self.ts_path, "Touchstone (*.s*p);;모든 파일 (*)")):
            row = QHBoxLayout()
            row.addWidget(QLabel(label))
            row.addWidget(field, 1)
            row.addWidget(button("찾기…", lambda checked=False, f=field, filt=filters: self.pick_file(f, filt)))
            layout.addLayout(row)
        self.load_button = button("파일 읽기", self.load_inputs, True)
        layout.addWidget(self.load_button, alignment=Qt.AlignmentFlag.AlignLeft)
        layout.addWidget(section_title("보고서 정보"))
        form = QFormLayout()
        self.title_input = QLineEdit(self.config.settings.title)
        self.product_input = QLineEdit()
        self.board_type = QComboBox()
        self.board_type.addItems(["MLO", "PCB"])
        self.date_input = QLineEdit(self.config.settings.report_date)
        self.revision_input = QLineEdit()
        self.ground_input = QLineEdit("GND, DGND")
        self.mirror = QCheckBox("Bottom을 하면에서 본 방향으로 좌우 반전")
        self.mirror.setChecked(True)
        for label, widget in (("제목", self.title_input), ("제품명", self.product_input), ("보드 종류", self.board_type), ("작성일", self.date_input), ("Revision", self.revision_input), ("접지 Net (쉼표 구분)", self.ground_input), ("Bottom 방향", self.mirror)):
            form.addRow(label, widget)
        layout.addLayout(form)
        self.input_summary = QLabel("파일이 없어도 채널·그룹에서 채널을 추가한 뒤 DCR 레포트를 만들 수 있습니다.")
        self.input_summary.setWordWrap(True)
        layout.addWidget(self.input_summary)
        layout.addStretch()
        return page

    def mapping_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        note = QLabel("사용할 포트를 선택하고 사이트·채널·Net을 확인하세요. 그룹이 같으면 한 그래프에 겹쳐 표시합니다. 빈 그룹은 상세 그래프를 제외합니다.")
        note.setWordWrap(True)
        layout.addWidget(note)
        row = QHBoxLayout()
        self.map_filter = QLineEdit()
        self.map_filter.setPlaceholderText("포트·Net·사이트·채널 검색")
        self.map_filter.textChanged.connect(self.filter_mapping)
        row.addWidget(self.map_filter, 1)
        row.addWidget(button("채널 추가", self.add_mapping))
        row.addWidget(button("선택 행 삭제", self.delete_mapping))
        layout.addLayout(row)
        row = QHBoxLayout()
        row.addWidget(button("검색 결과 사용", lambda: self.bulk_map(0, True)))
        row.addWidget(button("검색 결과 제외", lambda: self.bulk_map(0, False)))
        row.addWidget(button("선택 행 확인", lambda: self.bulk_map(2, True, selected_only=True)))
        self.group_input = QLineEdit()
        self.group_input.setPlaceholderText("선택 행에 적용할 그룹명 (빈 값 허용)")
        row.addWidget(self.group_input, 1)
        row.addWidget(button("그룹 적용", self.apply_group))
        layout.addLayout(row)
        self.mapping_table = make_table(["사용", "요약", "확인", "번호", "원본 포트", "SPD 포트", "Net", "사이트", "채널", "그룹", "Priority", "상태"])
        for index in range(4):
            self.mapping_table.setColumnWidth(index, 55)
        for index in (4, 5, 6, 8, 9):
            self.mapping_table.setColumnWidth(index, 180)
        self.mapping_table.itemChanged.connect(self.mapping_edited)
        layout.addWidget(self.mapping_table, 1)
        return page

    def cap_dcr_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        tabs = QTabWidget()
        cap = QWidget()
        cap_layout = QVBoxLayout(cap)
        label = QLabel("Enabled와 유효 모델이 있는 Automatic을 집계합니다. Disabled는 제외합니다. 수정값 예: 100 nF, 4.7 uF. 빈 수정값은 자동 판독으로 되돌립니다.")
        label.setWordWrap(True)
        cap_layout.addWidget(label)
        self.cap_tabs = QTabWidget()
        self.model_table = make_table(["모델", "이름에서 판독", "모델 주석에서 판독", "수정 용량", "근거·상태"])
        self.model_table.itemChanged.connect(self.model_edited)
        self.cap_tabs.addTab(self.model_table, "모델별 일괄 수정")
        instances = QWidget()
        instance_layout = QVBoxLayout(instances)
        self.cap_filter = QLineEdit()
        self.cap_filter.setPlaceholderText("RefDes 또는 모델 검색 · 최대 500개 표시")
        self.cap_filter.textChanged.connect(self.fill_components)
        instance_layout.addWidget(self.cap_filter)
        self.component_count = QLabel()
        instance_layout.addWidget(self.component_count)
        self.component_table = make_table(["RefDes", "모델", "상태", "Net", "현재 용량", "개별 수정"])
        self.component_table.itemChanged.connect(self.component_edited)
        instance_layout.addWidget(self.component_table)
        self.cap_tabs.addTab(instances, "부품별 수정")
        cap_layout.addWidget(self.cap_tabs)
        tabs.addTab(cap, "CAP")
        dcr = QWidget()
        dcr_layout = QVBoxLayout(dcr)
        dcr_note = QLabel("한 열을 가져오면 선택한 채널에 순서대로 대응합니다. 아래 표에서 사이트·채널·값을 확인하고 수정하세요. 표시값은 항상 mΩ입니다.")
        dcr_note.setWordWrap(True)
        dcr_layout.addWidget(dcr_note)
        row = QHBoxLayout()
        self.dcr_unit = QComboBox()
        self.dcr_unit.addItems(["mΩ", "Ω", "µΩ"])
        self.dcr_order = QComboBox()
        self.dcr_order.addItems(["사이트 → 채널", "채널 → 사이트"])
        row.addWidget(QLabel("입력 단위"))
        row.addWidget(self.dcr_unit)
        row.addWidget(self.dcr_order)
        row.addWidget(button("한 열 붙여넣기", self.paste_dcr))
        row.addWidget(button("Excel 범위 가져오기", self.import_dcr))
        row.addStretch()
        dcr_layout.addLayout(row)
        row = QHBoxLayout()
        row.addWidget(button("행 추가", self.add_dcr))
        row.addWidget(button("선택 행 삭제", self.delete_dcr))
        row.addStretch()
        dcr_layout.addLayout(row)
        self.dcr_table = make_table(["사이트", "채널", "DCR (mΩ)"])
        dcr_layout.addWidget(self.dcr_table, 1)
        tabs.addTab(dcr, "DCR")
        layout.addWidget(tabs)
        return page

    def preview_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        row = QHBoxLayout()
        row.addWidget(button("미리보기 갱신", self.preview))
        row.addWidget(button("Excel 저장…", self.export, True))
        row.addWidget(button("현재 그림 저장…", self.save_image))
        row.addStretch()
        layout.addLayout(row)
        self.page_summary = QLabel("입력과 연결을 확인한 뒤 미리보기를 갱신하세요.")
        self.page_summary.setWordWrap(True)
        layout.addWidget(self.page_summary)
        self.preview_selector = QComboBox()
        self.preview_selector.currentIndexChanged.connect(self.show_preview)
        layout.addWidget(self.preview_selector)
        self.preview_tabs = QTabWidget()
        self.preview_image = QLabel("도면과 그래프를 여기에 표시합니다.")
        self.preview_image.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.preview_image.setMinimumHeight(260)
        self.preview_image.setStyleSheet("background: #15181d;")
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(self.preview_image)
        self.preview_tabs.addTab(scroll, "그림")
        self.preview_table = QTextBrowser()
        self.preview_tabs.addTab(self.preview_table, "표 · 앞 40행")
        layout.addWidget(self.preview_tabs, 1)
        return page

    def error(self, error):
        self.log.appendPlainText(str(error))
        QMessageBox.warning(self, "확인할 내용", str(error))

    def pick_file(self, field, filters):
        path, _ = QFileDialog.getOpenFileName(self, "입력 파일 선택", field.text(), filters)
        if path:
            field.setText(path)

    def load_inputs(self):
        if self.runner.busy:
            return
        paths = {"spd": self.spd_path.text().strip(), "touchstone": self.ts_path.text().strip()}
        if not any(paths.values()):
            self.error("파일을 선택하세요. DCR만 입력할 때는 파일 읽기를 생략하세요.")
            return
        if self.pending_config is None and (self.spd or self.ts) and not self.confirm_unsaved():
            return
        self.action = "load"
        try:
            self.runner.start("load", paths)
        except Exception as exc:
            self.error(exc)

    def collect_config(self, *, dcr=None):
        config = self.config if dcr is None else replace(self.config, dcr=dcr)
        config.settings = ReportSettings(self.title_input.text().strip(), self.product_input.text().strip(), self.board_type.currentText(), self.date_input.text().strip(), self.revision_input.text().strip(), self.mirror.isChecked(), [n.strip() for n in self.ground_input.text().split(",") if n.strip()])
        rows = []
        for row in range(self.mapping_table.rowCount()):
            text = lambda col: self.mapping_table.item(row, col).text().strip()
            checked = lambda col: self.mapping_table.item(row, col).checkState() == Qt.CheckState.Checked
            rows.append(ChannelMapping(int(text(3)) if text(3) else None, text(4), text(5), text(6), text(7), text(8), checked(0), checked(1), text(9), text(10), text(11), checked(2)))
        config.mappings = rows
        if dcr is None:
            config.dcr = []
            for row in range(self.dcr_table.rowCount()):
                values = [self.dcr_table.item(row, col).text().strip() if self.dcr_table.item(row, col) else "" for col in range(3)]
                if not any(values):
                    continue
                try:
                    config.dcr.append(DcrRow(values[0], values[1], float(values[2].replace(",", ""))))
                except ValueError as exc:
                    raise ValueError(f"DCR {row+1}행 값을 확인하세요.") from exc
        validate_config(config)
        return config

    def fill_mappings(self):
        self.loading_table = True
        self.mapping_table.setRowCount(len(self.config.mappings))
        for row, mapping in enumerate(self.config.mappings):
            for col, value in enumerate((mapping.selected, mapping.summary, mapping.confirmed)):
                self.mapping_table.setItem(row, col, check_item(value))
            values = (mapping.port_index, mapping.source_name, mapping.spd_port, mapping.net, mapping.site, mapping.channel, mapping.group, mapping.priority, mapping.status)
            for col, value in enumerate(values, 3):
                self.mapping_table.setItem(row, col, text_item(value, col not in (3, 4, 11)))
        self.loading_table = False
        self.filter_mapping()

    def mapping_edited(self, item):
        if self.loading_table:
            return
        if item.column() == 2:
            row = item.row()
            name = self.mapping_table.item(row, 5).text().strip()
            if item.checkState() == Qt.CheckState.Checked and self.spd and name and not any(p.name == name for p in self.spd.ports):
                with QSignalBlocker(self.mapping_table):
                    item.setCheckState(Qt.CheckState.Unchecked)
                self.mapping_table.item(row, 11).setText("SPD 포트를 찾을 수 없습니다. 전체 이름을 다시 입력하세요.")
            else:
                self.mapping_table.item(row, 11).setText("사용자 확인" if item.checkState() == Qt.CheckState.Checked else "연결 확인 필요")
            return
        if item.column() not in (5, 6, 7, 8):
            return
        self.loading_table = True
        row = item.row()
        message = "연결이 바뀌었습니다. Net·사이트·채널을 검토하고 확인을 체크하세요."
        if item.column() == 5 and self.spd:
            name = item.text().strip()
            port = next((p for p in self.spd.ports if p.name == name), None)
            if port:
                self.mapping_table.item(row, 6).setText(port.net)
            elif name:
                self.mapping_table.item(row, 6).setText("")
                message = "SPD 포트를 찾을 수 없습니다. 전체 이름을 다시 입력하세요."
        self.mapping_table.item(row, 2).setCheckState(Qt.CheckState.Unchecked)
        self.mapping_table.item(row, 11).setText(message)
        self.loading_table = False

    def filter_mapping(self):
        needle = self.map_filter.text().casefold()
        for row in range(self.mapping_table.rowCount()):
            text = " ".join(self.mapping_table.item(row, col).text() for col in range(4, 12))
            self.mapping_table.setRowHidden(row, needle not in text.casefold())

    def selected_rows(self, table):
        return sorted({index.row() for index in table.selectionModel().selectedRows()})

    def bulk_map(self, column, state, selected_only=False):
        rows = self.selected_rows(self.mapping_table) if selected_only else range(self.mapping_table.rowCount())
        for row in rows:
            if not self.mapping_table.isRowHidden(row):
                self.mapping_table.item(row, column).setCheckState(Qt.CheckState.Checked if state else Qt.CheckState.Unchecked)
                if column == 2 and self.mapping_table.item(row, column).checkState() == Qt.CheckState.Checked:
                    self.mapping_table.item(row, 11).setText("사용자 확인")

    def apply_group(self):
        for row in self.selected_rows(self.mapping_table):
            if not self.mapping_table.isRowHidden(row):
                self.mapping_table.item(row, 9).setText(self.group_input.text().strip())

    def add_mapping(self):
        try:
            self.collect_config()
            self.config.mappings.append(ChannelMapping(site="SITE0", channel="", group=""))
            self.fill_mappings()
            self.mapping_table.scrollToBottom()
        except ValueError as exc:
            self.error(exc)

    def delete_mapping(self):
        for row in reversed(self.selected_rows(self.mapping_table)):
            self.mapping_table.removeRow(row)

    def fill_caps(self):
        from ..core.capacitance import format_capacitance
        self.loading_table = True
        cap_models = {c.model_name for c in self.spd.components if c.is_cap} if self.spd else set()
        models = [model for name, model in self.spd.models.items() if name in cap_models] if self.spd else []
        self.model_table.setRowCount(len(models))
        for row, model in enumerate(models):
            values = (model.name, format_capacitance(model.name_capacitance_f) if model.name_capacitance_f is not None else "미확인", format_capacitance(model.comment_capacitance_f) if model.comment_capacitance_f is not None else "미확인", format_capacitance(self.config.model_overrides[model.name]) if model.name in self.config.model_overrides else "", ("충돌 · " if model.conflict else "") + (model.evidence if model.valid else "빈 모델 / 미지정"))
            for col, value in enumerate(values):
                self.model_table.setItem(row, col, text_item(value, col == 3))
        self.loading_table = False
        self.model_table.resizeColumnsToContents()
        self.model_table.setColumnWidth(0, 230)
        self.fill_components()

    def fill_components(self):
        from ..core.capacitance import format_capacitance, resolved_capacitance
        self.loading_table = True
        needle = self.cap_filter.text().casefold()
        components = [c for c in self.spd.components if c.is_cap and needle in (c.refdes + " " + c.model_name).casefold()] if self.spd else []
        self.component_count.setText(f"검색 결과 {len(components):,}개 · {min(500, len(components)):,}개 표시")
        self.component_table.setRowCount(min(500, len(components)))
        for row, component in enumerate(components[:500]):
            value = resolved_capacitance(component, self.spd, self.config)
            values = (component.refdes, component.model_name, component.status, ", ".join(component.nets), format_capacitance(value) if value is not None else "미확인", format_capacitance(self.config.component_overrides[component.refdes]) if component.refdes in self.config.component_overrides else "")
            for col, value in enumerate(values):
                self.component_table.setItem(row, col, text_item(value, col == 5))
        self.loading_table = False

    def update_cap_override(self, table, item, expected_column, overrides):
        if self.loading_table or item.column() != expected_column:
            return False
        from ..core.capacitance import parse_capacitance
        key = table.item(item.row(), 0).text()
        try:
            if item.text().strip():
                value = parse_capacitance(item.text())
                if value is None or value <= 0:
                    raise ValueError("100 nF 또는 4.7 uF처럼 양수 용량과 단위를 입력하세요.")
                overrides[key] = value
            else:
                overrides.pop(key, None)
            return True
        except ValueError as exc:
            self.error(exc)
            self.fill_caps()
            return False

    def model_edited(self, item):
        self.update_cap_override(self.model_table, item, 3, self.config.model_overrides)
        if not self.loading_table:
            self.fill_components()

    def component_edited(self, item):
        if not self.update_cap_override(self.component_table, item, 5, self.config.component_overrides):
            return
        from ..core.capacitance import format_capacitance, resolved_capacitance
        row = item.row()
        refdes = self.component_table.item(row, 0).text()
        component = next(c for c in self.spd.components if c.refdes == refdes)
        value = resolved_capacitance(component, self.spd, self.config)
        with QSignalBlocker(self.component_table):
            self.component_table.item(row, 4).setText(format_capacitance(value) if value is not None else "미확인")

    def fill_dcr(self):
        self.dcr_table.setRowCount(len(self.config.dcr))
        for row, entry in enumerate(self.config.dcr):
            for col, value in enumerate((entry.site, entry.channel, f"{entry.resistance_mohm:.12g}")):
                self.dcr_table.setItem(row, col, text_item(value))

    def add_dcr(self):
        row = self.dcr_table.rowCount()
        self.dcr_table.insertRow(row)
        for col, value in enumerate(("SITE0", "", "")):
            self.dcr_table.setItem(row, col, text_item(value))

    def delete_dcr(self):
        for row in reversed(self.selected_rows(self.dcr_table)):
            self.dcr_table.removeRow(row)

    def confirm_dcr(self, values):
        config = self.collect_config(dcr=[])
        rows = assign_column(values, config.mappings, "site" if self.dcr_order.currentIndex() == 0 else "channel")
        config.dcr = rows
        validate_config(config)
        dialog = QDialog(self)
        dialog.setWindowTitle("DCR 대상과 순서 확인")
        dialog.resize(650, 500)
        layout = QVBoxLayout(dialog)
        layout.addWidget(QLabel("확인하면 기존 DCR 표를 이 값으로 교체합니다."))
        table = make_table(["사이트", "채널", "DCR (mΩ)"])
        table.setRowCount(len(rows))
        for index, row in enumerate(rows):
            for col, value in enumerate((row.site, row.channel, row.resistance_mohm)):
                table.setItem(index, col, text_item(value, False))
        layout.addWidget(table)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.config = config
            self.fill_dcr()
            return True
        return False

    def paste_dcr(self):
        try:
            if self.confirm_dcr(parse_column(QApplication.clipboard().text(), self.dcr_unit.currentText())):
                self.config.dcr_source_path = ""
        except ValueError as exc:
            self.error(exc)

    def import_dcr(self):
        path, _ = QFileDialog.getOpenFileName(self, "DCR Excel 선택", "", "Excel (*.xlsx)")
        if not path:
            return
        try:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True)
            names = book.sheetnames
            book.close()
            sheet, ok = QInputDialog.getItem(self, "DCR 시트", "시트", names, 0, False)
            if not ok:
                return
            cells, ok = QInputDialog.getText(self, "DCR 범위", "숫자만 있는 범위 (행 우선으로 읽음)", text="B2:B10")
            if ok:
                if self.confirm_dcr(read_range(path, sheet, cells.strip(), self.dcr_unit.currentText())):
                    self.config.dcr_source_path = str(Path(path).resolve())
        except Exception as exc:
            self.error(exc)

    def start_report(self, action, path=None):
        if self.runner.busy:
            return
        try:
            chosen = {"spd": self.spd_path.text().strip(), "touchstone": self.ts_path.text().strip()}
            if chosen != self.loaded_paths:
                raise ValueError("입력 경로가 바뀌었습니다. 먼저 파일 읽기를 눌러 새 입력을 적용하세요.")
            config = self.collect_config()
            from ..core.mapping import validate_mappings
            if any(m.selected for m in config.mappings):
                validate_mappings(config, self.ts, self.spd)
            self.action = action
            self.runner.start(action, {"config": config, "spd": self.spd, "ts": self.ts, "path": path, "folder": self.temp.name})
        except Exception as exc:
            self.error(exc)

    def preview(self):
        self.start_report("preview")

    def export(self):
        path, _ = QFileDialog.getSaveFileName(self, "Excel 레포트 저장", "PI_Report.xlsx", "Excel (*.xlsx)")
        if path:
            self.start_report("export", path if path.lower().endswith(".xlsx") else path + ".xlsx")

    def on_busy(self, busy):
        self.tabs.setEnabled(not busy)
        self.menuBar().setEnabled(not busy)
        self.cancel_button.setEnabled(busy)
        if busy:
            self.progress.setRange(0, 0)
            self.status.setText("처리 중입니다. 취소할 수 있습니다.")
        else:
            self.progress.setRange(0, 100)

    def on_progress(self, message, current, total):
        self.status.setText(message)
        if total:
            self.progress.setRange(0, 100)
            self.progress.setValue(int(current / total * 100))

    def on_failed(self, error):
        self.pending_config = None
        self.status.setText("작업을 완료하지 못했습니다. 아래 내용을 확인하세요.")
        self.error(error)

    def on_finished(self, result):
        self.progress.setValue(100)
        if self.action == "load":
            self.clear_preview()
            self.loaded_paths = {"spd": self.spd_path.text().strip(), "touchstone": self.ts_path.text().strip()}
            before = [(v.path, v.signature) if v else None for v in (self.spd, self.ts)]
            self.spd, self.ts, mappings, errors = result
            after = [(v.path, v.signature) if v else None for v in (self.spd, self.ts)]
            if self.pending_config is not None:
                self.config = self.pending_config
                self.pending_config = None
            elif before != after:
                self.config.mappings = mappings
                self.config.model_overrides.clear()
                self.config.component_overrides.clear()
                if any(before) and self.config.dcr:
                    self.config.dcr.clear()
                    self.log.appendPlainText("입력 파일이 바뀌어 이전 DCR 값을 비웠습니다. 새 입력에 맞는 DCR을 가져오세요.")
            self.fill_mappings()
            self.fill_caps()
            self.fill_dcr()
            summary = f"SPD 부품 {len(self.spd.components) if self.spd else 0:,}개 · Touchstone 포트 {len(self.ts.ports) if self.ts else 0:,}개 · 채널 연결 {len(mappings):,}개"
            self.input_summary.setText(summary)
            self.log.appendPlainText(summary)
            for warning in errors + (self.spd.warnings if self.spd else []) + (self.ts.warnings if self.ts else []):
                self.log.appendPlainText(warning)
            self.tabs.setCurrentIndex(1)
            self.status.setText("파일을 읽었습니다. 사용할 포트와 채널 연결을 확인하세요.")
        else:
            if self.action == "preview":
                report, self.preview_images = result
                self.preview_book = report.path
                self.preview_selector.clear()
                for name, path in self.preview_images:
                    self.preview_selector.addItem(f"그림 · {name}", ("image", path))
                for sheet in report.sheets:
                    self.preview_selector.addItem(f"표 · {sheet}", ("table", sheet))
                self.show_preview()
                self.status.setText("미리보기를 갱신했습니다. 설정을 바꾸면 다시 갱신하세요.")
            else:
                report = result
                self.status.setText(f"저장 완료: {report.path}")
                QMessageBox.information(self, "저장 완료", f"{len(report.sheets)}개 페이지를 저장했습니다.\n{report.path}")
            self.page_summary.setText("포함: " + ", ".join(report.sheets))
            self.log.appendPlainText("포함 페이지: " + ", ".join(report.sheets))
            for name, reason in report.omitted.items():
                self.log.appendPlainText(f"제외 · {name}: {reason}")
            for warning in report.warnings:
                self.log.appendPlainText(warning)

    def show_preview(self):
        data = self.preview_selector.currentData()
        if not data:
            return
        kind, name = data
        if kind == "image":
            pixmap = QPixmap(name)
            self.preview_image.setPixmap(pixmap.scaled(max(400, self.preview_tabs.width()-35), max(300, self.preview_tabs.height()-65), Qt.AspectRatioMode.KeepAspectRatio, Qt.TransformationMode.SmoothTransformation))
            self.preview_tabs.setCurrentIndex(0)
        elif self.preview_book:
            from openpyxl import load_workbook
            book = load_workbook(self.preview_book, read_only=True, data_only=True)
            try:
                sheet = book[name]
                rows = []
                for row in sheet.iter_rows(min_row=1, max_row=min(40, sheet.max_row), max_col=min(40, sheet.max_column), values_only=True):
                    cells = "".join(f"<td>{html.escape(str(v) if v is not None else '')}</td>" for v in row)
                    rows.append(f"<tr>{cells}</tr>")
                self.preview_table.setHtml("<style>table{border-collapse:collapse;background:white;color:#222}td{border:1px solid #ccc;padding:5px;white-space:nowrap}</style><table>" + "".join(rows) + "</table>")
            finally:
                book.close()
            self.preview_tabs.setCurrentIndex(1)

    def clear_preview(self):
        self.preview_book = None
        self.preview_images = []
        self.preview_selector.clear()
        self.preview_image.clear()
        self.preview_table.clear()
        self.page_summary.setText("입력과 연결을 확인한 뒤 미리보기를 갱신하세요.")

    def save_image(self):
        data = self.preview_selector.currentData()
        if not data or data[0] != "image":
            self.error("미리보기에서 저장할 그림을 먼저 선택하세요.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "그림 저장", "placement.png", "PNG (*.png)")
        if path:
            try:
                shutil.copyfile(data[1], path if path.lower().endswith(".png") else path + ".png")
            except OSError as exc:
                self.error(exc)

    def apply_settings(self):
        settings = self.config.settings
        self.title_input.setText(settings.title)
        self.product_input.setText(settings.product)
        self.board_type.setCurrentText(settings.board_type)
        self.date_input.setText(settings.report_date)
        self.revision_input.setText(settings.revision)
        self.ground_input.setText(", ".join(settings.ground_nets))
        self.mirror.setChecked(settings.bottom_mirror)

    def save_project(self):
        if self.runner.busy:
            return False
        path, _ = QFileDialog.getSaveFileName(self, "프로젝트 저장", self.project_path or "project.report.json", "프로젝트 (*.report.json)")
        if path:
            try:
                path = path if path.lower().endswith(".report.json") else path + ".report.json"
                sources = {"spd": self.spd_path.text().strip(), "touchstone": self.ts_path.text().strip()}
                signatures = {key: source.signature for key, source in (("spd", self.spd), ("touchstone", self.ts)) if source and sources[key] and Path(sources[key]).resolve() == Path(source.path).resolve()}
                save_project(path, self.collect_config(), sources, signatures)
                self.project_path = path
                self.status.setText(f"프로젝트 저장 완료: {path}")
                self.saved_state = self.state_text()
                return True
            except Exception as exc:
                self.error(exc)
        return False

    def state_text(self):
        return json.dumps({"config": asdict(self.collect_config()), "spd": self.spd_path.text(), "touchstone": self.ts_path.text()}, ensure_ascii=False, sort_keys=True)

    def confirm_unsaved(self):
        try:
            if self.state_text() == self.saved_state:
                return True
        except ValueError:
            pass
        choice = QMessageBox.question(self, "저장하지 않은 설정", "현재 프로젝트 설정을 저장할까요?", QMessageBox.StandardButton.Save | QMessageBox.StandardButton.Discard | QMessageBox.StandardButton.Cancel)
        if choice == QMessageBox.StandardButton.Save:
            return self.save_project()
        return choice == QMessageBox.StandardButton.Discard

    def open_project(self):
        if self.runner.busy:
            return
        path, _ = QFileDialog.getOpenFileName(self, "프로젝트 열기", "", "프로젝트 (*.report.json)")
        if not path:
            return
        if not self.confirm_unsaved():
            return
        try:
            config, sources, _, changed = load_project(path)
            self.clear_preview()
            self.config = config
            self.project_path = path
            self.apply_settings()
            self.fill_mappings()
            self.fill_dcr()
            self.spd_path.setText(sources.get("spd", ""))
            self.ts_path.setText(sources.get("touchstone", ""))
            self.saved_state = self.state_text()
            self.spd = self.ts = None
            self.loaded_paths = {"spd": "", "touchstone": ""}
            self.fill_caps()
            if changed:
                self.config.dcr.clear()
                self.fill_dcr()
                self.log.appendPlainText("원본이 변경되었거나 없습니다. 파일을 다시 읽고 채널 연결과 CAP 수정값을 확인하세요.")
                self.log.appendPlainText("이전 원본에 연결된 DCR은 비웠습니다. 필요한 값은 다시 가져오세요.")
            if any(sources.values()):
                self.pending_config = config
                self.load_inputs()
            else:
                self.status.setText("프로젝트를 열었습니다.")
        except Exception as exc:
            self.error(exc)

    def new_project(self):
        if self.runner.busy:
            return
        if not self.confirm_unsaved():
            return
        self.config = ReportConfig()
        self.spd = self.ts = None
        self.project_path = None
        self.loaded_paths = {"spd": "", "touchstone": ""}
        self.spd_path.clear()
        self.ts_path.clear()
        self.apply_settings()
        self.fill_mappings()
        self.fill_caps()
        self.fill_dcr()
        self.clear_preview()
        self.log.clear()
        self.tabs.setCurrentIndex(0)
        self.saved_state = self.state_text()

    def open_help(self):
        path = Path(__file__).resolve().parents[1] / "docs" / "help.html"
        if not QDesktopServices.openUrl(QUrl.fromLocalFile(str(path))):
            self.error(f"설명서를 열 수 없습니다: {path}")

    def closeEvent(self, event):
        if self.runner.busy:
            if QMessageBox.question(self, "종료", "진행 중인 작업을 취소하고 종료할까요?") != QMessageBox.StandardButton.Yes:
                event.ignore()
                return
            self.runner.cleanup(force=True)
        if not self.confirm_unsaved():
            event.ignore()
            return
        self.temp.cleanup()
        event.accept()
