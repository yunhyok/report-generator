"""Headless checks using the same parsers, mappings and report writer as the GUI."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import asdict
import html
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from time import perf_counter

from openpyxl import load_workbook

for _variable in ("OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "OMP_NUM_THREADS"):
    os.environ.setdefault(_variable, "1")

from . import __version__
from .core.capacitance import parse_capacitance, summarize_caps
from .core.mapping import validate_mappings
from .core.models import CancelledError, checkpoint, source_signature
from .core.project import load_project, validate_config
from .core.report import FREQUENCIES, FREQUENCY_LABELS, build_report
from .core.spd import read_spd
from .core.touchstone import interpolate_impedance, read_touchstone

RELATIVE_TOLERANCE = 1e-8
ABSOLUTE_TOLERANCE = 1e-12


def _number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _site(value):
    text = str(int(value)) if _number(value) and value == int(value) else str(value or "").strip()
    match = re.fullmatch(r"(?:SITE)?(\d+)", text, re.I)
    return f"SITE{int(match[1])}" if match else text


def _frequency(value):
    text = str(value or "").split(" [", 1)[0].strip().casefold()
    return next((label for label in FREQUENCY_LABELS if label.casefold() == text), None)


def _scale(value):
    return .001 if re.search(r"\[\s*m[ΩΩ]\s*\]", str(value)) else 1.


def _nominal(value):
    text = str(value or "")
    if "unknown" in text.casefold() or "미확인" in text:
        return "Unknown"
    capacitance = parse_capacitance(text)
    return f"{capacitance:.12g}F" if capacitance is not None else None


def _put(records, key, cell, scale=1.):
    if _number(cell.value):
        records[key] = {"value": float(cell.value) * scale, "sheet": cell.parent.title, "cell": cell.coordinate}


def _pi_tables(sheet):
    result = {}
    for header in sheet.iter_rows():
        for site_cell in header:
            if str(site_cell.value).upper() != "SITE":
                continue
            col, row = site_cell.column, site_cell.row
            frequencies = [(c.column, _frequency(c.value)) for c in header[col:col + 6] if _frequency(c.value)]
            if not frequencies:
                continue
            above = sheet.cell(row - 1, col).value if row > 1 else None
            fixed = above.split(" / [", 1)[0].strip() if isinstance(above, str) and " / [" in above and not above.upper().startswith("CHANNEL") else None
            channel_col = next((c.column for c in reversed(header[:col - 1]) if isinstance(c.value, str) and c.value.upper().startswith("CHANNEL")), None)
            if not fixed and channel_col is None:
                continue
            unit = above if fixed else sheet.cell(row, channel_col).value
            if "[" not in str(unit):
                unit = sheet.cell(row, frequencies[0][0]).value
            channel = fixed
            for r in range(row + 1, sheet.max_row + 1):
                raw_site = sheet.cell(r, col).value
                if raw_site is None or str(raw_site).upper() == "SITE":
                    break
                if not fixed and isinstance(sheet.cell(r, channel_col).value, str):
                    channel = sheet.cell(r, channel_col).value.strip()
                if channel:
                    for c, label in frequencies:
                        _put(result, (channel, _site(raw_site), label), sheet.cell(r, c), _scale(unit))
    return result


def _matrix_tables(sheet, cap=False, site_names=()):
    result = {}
    headers = []
    for row in sheet.iter_rows():
        sites = [(c.column, _site(c.value)) for c in row if isinstance(c.value, str) and
                 (re.fullmatch(r"SITE\d+", c.value, re.I) or _site(c.value) in site_names)]
        if sites:
            channel = next((c.column for c in row if isinstance(c.value, str) and c.value.upper().startswith("CHANNEL")), sites[0][0] - 1)
            headers.append((row[0].row, sites, channel))
    for index, (row, sites, col) in enumerate(headers):
        label = sheet.cell(row, col).value
        nominal = None
        if cap:
            candidates = [label] + [sheet.cell(max(1, row - 1), c).value for c in range(col, sites[-1][0] + 1)]
            nominal = next((value for candidate in candidates if (value := _nominal(candidate)) is not None), None)
            if nominal is None:
                continue
        end = headers[index + 1][0] if index + 1 < len(headers) else sheet.max_row + 1
        for r in range(row + 1, end):
            channel = sheet.cell(r, col).value
            if not isinstance(channel, str) or not channel.strip():
                continue
            for c, site in sites:
                key = (channel.strip(), site, nominal) if cap else (channel.strip(), site)
                _put(result, key, sheet.cell(r, c), 1. if cap else _scale(label))
    # Zero count cells are table padding, not additional capacitor groups.
    return {key: item for key, item in result.items() if not cap or item["value"] != 0}


def inspect_workbook(path, site_names=()):
    """Align report cells by channel/site/frequency/nominal, never by row number."""
    book = load_workbook(path, data_only=True)
    values = {kind: {} for kind in ("PI", "PI100", "PI_DETAIL", "CAP", "DCR")}
    layout = []
    try:
        for sheet in book:
            normalized = sheet.title.replace("_", " ").upper()
            pi = _pi_tables(sheet)
            if pi:
                values["PI" if normalized == "PI ALL" else "PI_DETAIL"].update(pi)
            elif sheet.title == "CAP_수량":
                values["CAP"].update(_matrix_tables(sheet, cap=True, site_names=site_names))
            elif normalized == "DCR":
                values["DCR"].update(_matrix_tables(sheet, site_names=site_names))
            elif normalized == "PI ALL@100MHZ":
                values["PI100"].update({(*key, "100MHz"): item for key, item in _matrix_tables(sheet, site_names=site_names).items()})
            fonts, fills = Counter(), Counter()
            bold = numeric = populated = 0
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None and not cell.has_style:
                        continue
                    populated += cell.value is not None
                    numeric += _number(cell.value)
                    bold += bool(cell.font.bold)
                    fonts[f"{cell.font.name or 'default'} / {cell.font.sz or 'default'} / {'bold' if cell.font.bold else 'regular'}"] += 1
                    color = cell.fill.fgColor
                    fills[f"{cell.fill.fill_type or 'none'} / {color.type}:{color.value}"] += 1
            layout.append({"sheet": sheet.title, "dimension": sheet.calculate_dimension(),
                           "populated_cells": populated, "numeric_cells": numeric, "bold_cells": bold,
                           "fonts": dict(fonts), "fills": dict(fills), "freeze_panes": sheet.freeze_panes,
                           "selection_panes": [s.pane for s in sheet.sheet_view.selection if s.pane],
                           "images": len(sheet._images), "gridlines": sheet.sheet_view.showGridLines})
        return values, layout
    finally:
        book.close()


def compare_values(actual, expected):
    common = actual.keys() & expected.keys()
    differences = []
    absolute = []
    for key in sorted(common):
        a, e = actual[key]["value"], expected[key]["value"]
        absolute.append(abs(a - e))
        if not math.isclose(a, e, rel_tol=RELATIVE_TOLERANCE, abs_tol=ABSOLUTE_TOLERANCE):
            differences.append({"key": list(key), "actual": actual[key], "expected": expected[key]})
    return {"actual_count": len(actual), "expected_count": len(expected), "compared": len(common),
            "matches": len(common) - len(differences), "differences": len(differences),
            "missing": len(expected.keys() - actual.keys()), "extra": len(actual.keys() - expected.keys()),
            "max_absolute_error": max(absolute, default=0.), "difference_examples": differences[:20],
            "missing_examples": [list(k) for k in sorted(expected.keys() - actual.keys())[:20]],
            "extra_examples": [list(k) for k in sorted(actual.keys() - expected.keys())[:20]]}


def _expected(config, spd, ts, omitted):
    values = {kind: {} for kind in ("PI", "PI100", "PI_DETAIL", "CAP", "DCR")}
    if ts:
        for mapping in config.mappings:
            if not (mapping.selected and mapping.confirmed and mapping.port_index is not None):
                continue
            for frequency, label in zip(FREQUENCIES, FREQUENCY_LABELS):
                value, _ = interpolate_impedance(ts, mapping.port_index, frequency)
                if value is None:
                    continue
                key, item = (mapping.channel, _site(mapping.site), label), {"value": value}
                if mapping.summary:
                    values["PI"][key] = item
                    if frequency == 1e8:
                        values["PI100"][key] = item
                if mapping.group and f"그룹: {mapping.group}" not in omitted:
                    values["PI_DETAIL"][key] = item
    if spd:
        for row in summarize_caps(spd, config):
            values["CAP"][row["channel"], _site(row["site"]), _nominal(row["label"])] = {"value": float(row["count"])}
    for row in config.dcr:
        values["DCR"][row.channel, _site(row.site)] = {"value": row.resistance_mohm * .001}
    return values


def _html(audit):
    def table(comparisons):
        rows = "".join(f"<tr><td>{html.escape(name)}</td>" + "".join(f"<td>{values[key]}</td>" for key in ("compared", "matches", "differences", "missing", "extra")) + "</tr>" for name, values in comparisons.items())
        return "<table><tr><th>항목</th><th>비교</th><th>일치</th><th>차이</th><th>누락</th><th>추가</th></tr>" + rows + "</table>"
    body = f"<h1>프로그램 계산 → 출력 셀: {html.escape(audit['status'])}</h1>"
    if "error" in audit:
        body += f"<p>{html.escape(audit['error'])}</p>"
    if "production_checks" in audit:
        body += "<h2>실제 프로그램 계산과 내보낸 셀</h2>" + table(audit["production_checks"])
    if "reference_comparison" in audit:
        body += f"<h2>참조 문서와 비교: {html.escape(audit['reference_status'])}</h2><p>참조와의 차이만으로 프로그램 오류를 뜻하지 않습니다. 입력 버전·선택 채널·모델 수정값을 함께 확인하세요.</p>" + table(audit["reference_comparison"])
    body += "<h2>상세 근거 · 서식 · 입력 정보</h2><pre>" + html.escape(json.dumps(audit, ensure_ascii=False, indent=2)) + "</pre>"
    return '<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>보고서 검증</title><style>body{background:#202226;color:#eceef1;font:16px/1.6 system-ui;margin:2rem}h1,h2{color:#8fc1ff}table{border-collapse:collapse}td,th{border:1px solid #66717e;padding:.4rem 1rem}th,pre{background:#2b2e34}pre{padding:1rem;white-space:pre-wrap;overflow-wrap:anywhere}</style>' + body + "</html>"


def _write_audits(audit, json_path, html_path):
    for path, text in ((json_path, json.dumps(audit, ensure_ascii=False, indent=2, allow_nan=False) + "\n"), (html_path, _html(audit))):
        temporary = None
        try:
            with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(text)
            os.replace(temporary, path)
        finally:
            if temporary:
                temporary.unlink(missing_ok=True)


def validate_project(project_path, output_path, reference_path=None, *, progress=None, cancel=None):
    """Generate with production code, then audit its numbers and optional reference."""
    started = perf_counter()
    project, output = Path(project_path).expanduser().resolve(), Path(output_path).expanduser().resolve()
    reference = Path(reference_path).expanduser().resolve() if reference_path else None
    json_path, html_path = output.with_suffix(".validation.json"), output.with_suffix(".validation.html")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("출력 파일 확장자는 .xlsx여야 합니다.")
    protected = {project, *([reference] if reference else [])}
    if any(path in protected for path in (output, json_path, html_path)):
        raise ValueError("입력·참조·프로젝트 파일에 검증 결과를 덮어쓸 수 없습니다.")
    audit = {"schema_version": 1, "application_version": __version__, "status": "error", "project": str(project),
             "output": str(output), "reference": str(reference) if reference else None,
             "tolerance": {"relative": RELATIVE_TOLERANCE, "absolute": ABSOLUTE_TOLERANCE},
             "units": {"PI": "ohm", "PI100": "ohm", "PI_DETAIL": "ohm", "DCR": "ohm", "CAP": "count"}}
    try:
        config, sources, signatures, changed = load_project(project)
    except Exception as exc:
        # An unreadable project cannot establish its input paths; only create new diagnostics.
        if not json_path.exists() and not html_path.exists():
            output.parent.mkdir(parents=True, exist_ok=True)
            audit.update(error=str(exc), elapsed_seconds=round(perf_counter() - started, 3))
            _write_audits(audit, json_path, html_path)
        raise
    protected.update(Path(source).resolve() for source in sources.values() if source)
    if config.dcr_source_path:
        protected.add((project.parent / config.dcr_source_path).resolve())
    if any(path in protected for path in (output, json_path, html_path)):
        raise ValueError("입력·참조·프로젝트 파일에 검증 결과를 덮어쓸 수 없습니다.")
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        checkpoint(cancel)
        if changed:
            raise ValueError("프로젝트 원본이 변경되었거나 없습니다. 자동 확인하지 않습니다: " + ", ".join(changed))
        validate_config(config)
        parsed, errors = {}, []
        for kind, reader in (("spd", read_spd), ("touchstone", read_touchstone)):
            if sources.get(kind):
                try:
                    parsed[kind] = reader(sources[kind], progress=progress, cancel=cancel)
                except CancelledError:
                    raise
                except Exception as exc:
                    errors.append(f"{kind}: {exc}")
        if any(sources.values()) and not parsed:
            raise ValueError("입력 파일을 읽을 수 없습니다: " + " | ".join(errors))
        spd, ts = parsed.get("spd"), parsed.get("touchstone")
        if any(m.selected for m in config.mappings):
            validate_mappings(config, ts, spd)
        for kind, source in sources.items():
            if source and source_signature(source) != signatures[kind]:
                raise ValueError("검증 중 원본이 변경되었습니다. 프로젝트 연결을 다시 확인하세요.")
        audit["sources"] = {kind: {"path": source, "signature": signatures[kind]} for kind, source in sources.items() if source}
        audit["input_errors"] = errors
        audit["selected_mappings"] = [asdict(m) for m in config.mappings if m.selected]
        with tempfile.TemporaryDirectory(prefix="report-validation-", dir=output.parent) as folder:
            staged = Path(folder) / output.name
            result = build_report(config, spd, ts, staged, progress=progress, cancel=cancel)
            site_names = {_site(row.site) for row in [*config.mappings, *config.dcr]}
            actual, layout = inspect_workbook(staged, site_names)
            expected = _expected(config, spd, ts, result.omitted)
            audit["production_checks"] = {kind: compare_values(actual[kind], expected[kind]) for kind in actual}
            audit["layout"] = layout
            audit["sheets"], audit["omitted"], audit["warnings"] = result.sheets, result.omitted, result.warnings
            if reference:
                reference_values, reference_layout = inspect_workbook(reference, site_names)
                audit["reference_comparison"] = {kind: compare_values(actual[kind], reference_values[kind]) for kind in actual}
                comparisons = list(audit["reference_comparison"].values())
                audit["reference_status"] = ("not_comparable" if not any(check["compared"] for check in comparisons) else
                                             "differences" if any(check["differences"] or check["missing"] or check["extra"] for check in comparisons) else "matches")
                audit["reference_layout"] = reference_layout
                audit["reference_note"] = "Differences do not by themselves establish an application defect; verify source revisions and selected mappings."
            checks = audit["production_checks"].values()
            audit["status"] = "passed" if all(not (check["differences"] or check["missing"] or check["extra"]) for check in checks) else "production_difference"
            checkpoint(cancel)
            for kind, source in sources.items():
                if source and source_signature(source) != signatures[kind]:
                    raise ValueError("검증 중 원본이 변경되었습니다. 파일을 다시 읽으세요.")
            os.replace(staged, output)
        audit["elapsed_seconds"] = round(perf_counter() - started, 3)
        _write_audits(audit, json_path, html_path)
        return audit
    except (Exception, KeyboardInterrupt) as exc:
        audit.update(status="cancelled" if isinstance(exc, (KeyboardInterrupt, CancelledError)) else "error",
                     error="작업을 취소했습니다." if isinstance(exc, KeyboardInterrupt) else str(exc),
                     elapsed_seconds=round(perf_counter() - started, 3))
        _write_audits(audit, json_path, html_path)
        raise


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate-project", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--reference", type=Path)
    args = parser.parse_args(argv)
    try:
        audit = validate_project(args.validate_project, args.output, args.reference)
        if sys.stdout is not None:
            print(json.dumps({"status": audit["status"], "reference_status": audit.get("reference_status"), "output": audit["output"]}, ensure_ascii=False))
        return 0 if audit["status"] == "passed" else 1
    except KeyboardInterrupt:
        return 130
    except Exception as exc:
        if sys.stderr is not None:
            print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
