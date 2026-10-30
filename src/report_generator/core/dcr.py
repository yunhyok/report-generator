from __future__ import annotations

import math
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .models import ChannelMapping, DcrRow


def natural_key(value: str) -> list:
    return [int(p) if p.isdigit() else p.casefold() for p in re.split(r"(\d+)", value)]


def numeric_value(value, unit: str = "mΩ") -> float:
    if unit not in ("mΩ", "Ω", "µΩ"):
        raise ValueError("DCR 단위를 확인하세요.")
    if isinstance(value, bool) or value is None:
        raise ValueError("DCR 셀이 비어 있거나 숫자가 아닙니다.")
    try:
        number = float(str(value).strip().replace(",", ""))
    except ValueError as exc:
        raise ValueError(f"DCR 숫자를 확인하세요: {value}") from exc
    if not math.isfinite(number) or number < 0:
        raise ValueError("DCR은 0 이상의 유한한 숫자여야 합니다.")
    return number * {"mΩ": 1, "Ω": 1000, "µΩ": 0.001}[unit]


def parse_column(text: str, unit: str = "mΩ") -> list[float]:
    lines = text.strip().splitlines()
    if not lines:
        raise ValueError("붙여넣을 값이 없습니다.")
    values = []
    for index, line in enumerate(lines, 1):
        if "\t" in line.strip():
            raise ValueError(f"{index}행: 한 열만 복사하세요.")
        values.append(numeric_value(line, unit))
    return values


def read_range(path: str | Path, sheet: str, cell_range: str, unit: str = "mΩ") -> list[float]:
    try:
        left, top, right, bottom = range_boundaries(cell_range)
    except (ValueError, TypeError) as exc:
        raise ValueError("A1:A10처럼 셀 범위를 지정하세요.") from exc
    if None in (left, top, right, bottom) or min(left, top) < 1 or right < left or bottom < top or (right-left+1)*(bottom-top+1) > 100_000:
        raise ValueError("범위는 100,000셀 이하의 유효한 영역이어야 합니다.")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in book.sheetnames:
            raise ValueError("시트 이름을 확인하세요.")
        values = []
        for row_index, row in enumerate(book[sheet].iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right), top):
            for column_index, cell in enumerate(row, left):
                try:
                    values.append(numeric_value(cell.value, unit))
                except ValueError as exc:
                    raise ValueError(f"{row_index}행 {column_index}열: {exc}") from exc
        if len(values) != (right-left+1)*(bottom-top+1):
            raise ValueError("선택한 범위에 빈 셀이 있습니다.")
        return values
    finally:
        book.close()


def ordered_targets(mappings: list[ChannelMapping], order: str = "site") -> list[tuple[str, str]]:
    keys = {(m.site.strip(), m.channel.strip()) for m in mappings if m.selected and m.site.strip() and m.channel.strip()}
    return sorted(keys, key=lambda key: (natural_key(key[0]), natural_key(key[1])) if order == "site" else (natural_key(key[1]), natural_key(key[0])))


def assign_column(values: list[float], mappings: list[ChannelMapping], order: str = "site") -> list[DcrRow]:
    keys = ordered_targets(mappings, order)
    if not keys:
        raise ValueError("채널·그룹에서 사이트와 채널을 먼저 추가하세요. DCR만 입력할 수도 있습니다.")
    if len(values) != len(keys):
        raise ValueError(f"입력 {len(values)}개, 선택 대상 {len(keys)}개입니다. 순서와 행 수를 확인하세요.")
    return [DcrRow(site, channel, value) for (site, channel), value in zip(keys, values)]
