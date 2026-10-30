from __future__ import annotations

from math import isfinite
import os
from pathlib import Path
import re
import tempfile

import numpy as np
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.views import Selection

from .models import ReportConfig, ReportResult, SpdData, TouchstoneData, checkpoint
from .render import mapped_nets, placement_components, render_group_plot, render_placement
from .capacitance import component_active, summarize_caps
from .touchstone import interpolate_impedance


FREQUENCIES = (1e5, 1e6, 1e7, 1e8, 1e9)
FREQUENCY_LABELS = ("100kHz", "1MHz", "10MHz", "100MHz", "1GHz")
WHITE = PatternFill("solid", fgColor="FFFFFF")
HEADER = PatternFill("solid", fgColor="E8EDF3")
LINE = Side(style="thin", color="B8BEC7")


def _mappings(config, ts=None, summary=False):
    return [m for m in config.mappings if m.selected and m.confirmed and
            (not summary or m.summary) and
            (ts is None or (m.port_index is not None and 1 <= m.port_index <= len(ts.ports)))]


def _spot_rows(config, ts, summary=False):
    if ts is None:
        return []
    rows = []
    for m in _mappings(config, ts, summary):
        values, interpolated = [], []
        for frequency in FREQUENCIES:
            value, interp = interpolate_impedance(ts, m.port_index, frequency)
            values.append(value)
            interpolated.append(interp)
        rows.append((m, values, interpolated))
    return rows


def _safe_sheet_name(label, used):
    label = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", str(label)).strip("'") or "Group"
    base = label[:31]
    candidate, number = base, 2
    while candidate.casefold() in used:
        suffix = f" ({number})"
        candidate = base[:31 - len(suffix)] + suffix
        number += 1
    used.add(candidate.casefold())
    return candidate


def _plan(config, spd, ts):
    pages, omitted, used = [], {}, set()
    spots = _spot_rows(config, ts, summary=True)
    def add(kind, label, payload=None):
        pages.append((kind, _safe_sheet_name(label, used), payload))
    if any(values[3] is not None for _, values, _ in spots):
        add("100", "PI ALL@100MHz", spots)
    else:
        omitted["PI ALL@100MHz"] = "선택한 포트에 유효한 100MHz 값이 없습니다."
    if any(any(value is not None for value in values) for _, values, _ in spots):
        add("pi", "PI ALL", spots)
    else:
        omitted["PI ALL"] = "선택한 포트에 5개 기준 주파수의 유효한 값이 없습니다."
    if config.dcr:
        add("dcr", "DCR", config.dcr)
    else:
        omitted["DCR"] = "DCR 입력이 없습니다."
    caps = summarize_caps(spd, config) if spd is not None else []
    if caps:
        add("cap", "CAP_수량", caps)
    else:
        omitted["CAP_수량"] = "선택한 채널/사이트의 활성 커패시터 수량이 없습니다."
    sides = []
    if spd is not None:
        for side in (("Top", "Bottom") if config.settings.board_type == "PCB" else ("Top",)):
            if placement_components(spd, config, side):
                sides.append(side)
    if sides:
        add("placement", "CAP_배치도면", sides)
    else:
        omitted["CAP_배치도면"] = "표시할 부품 좌표가 없습니다."
    groups = list(dict.fromkeys(m.group for m in _mappings(config, ts) if m.group)) if ts else []
    for group in groups:
        maps = [m for m in _mappings(config, ts) if m.group == group]
        valid = any(np.any(np.isfinite(np.abs(ts.z_diagonal[:, m.port_index - 1])) &
                           (np.abs(ts.z_diagonal[:, m.port_index - 1]) > 0) &
                           np.isfinite(ts.frequency_hz) & (ts.frequency_hz >= 1e3) & (ts.frequency_hz <= 1e9)) for m in maps)
        if valid:
            add("group", group, group)
        else:
            omitted[f"그룹: {group}"] = "1kHz~1GHz 범위에 유효한 임피던스가 없습니다."
    return pages, omitted


def plan_pages(config: ReportConfig, spd: SpdData | None = None,
               ts: TouchstoneData | None = None) -> tuple[list[str], dict[str, str]]:
    pages, omitted = _plan(config, spd, ts)
    return [name for _, name, _ in pages], omitted


def _put(ws, row, column, value, *, header=False, number_format=None):
    cell = ws.cell(row, column)
    if isinstance(value, str):
        # Explicit string type prevents spreadsheet formulas from source names.
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
        cell.data_type = "s"
    else:
        cell.value = value
    cell.font = Font(name="Malgun Gothic", size=11, bold=header, color="202226")
    cell.fill = HEADER if header else WHITE
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=header)
    cell.border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
    if number_format:
        cell.number_format = number_format
    return cell


def _table(ws, headers, rows, row=4, column=2, number_columns=()):
    for i, label in enumerate(headers):
        _put(ws, row, column + i, label, header=True)
    for offset, values in enumerate(rows, row + 1):
        for i, value in enumerate(values):
            _put(ws, offset, column + i, value,
                 number_format="0.000" if i in number_columns else None)
    ws.row_dimensions[row].height = 30


def _base_sheet(wb, name, config):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.freeze_panes = "C5"
    ws.column_dimensions["A"].width = 3
    for i in range(2, 60):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["P"].width = 27
    title = ws.cell(1, 2, ILLEGAL_CHARACTERS_RE.sub("", f"{config.settings.title} · {name}")[:32767])
    title.data_type = "s"
    title.font = Font(name="Malgun Gothic", size=24, bold=True, color="202226")
    title.fill = WHITE
    ws.row_dimensions[1].height = 38
    subtitle = " · ".join(str(v) for v in [config.settings.product, config.settings.board_type,
                          config.settings.report_date, config.settings.revision] if v)
    ws.cell(2, 2, ILLEGAL_CHARACTERS_RE.sub("", subtitle)[:32767]).data_type = "s"
    ws.cell(2, 2).font = Font(name="Malgun Gothic", size=11, color="555555")
    ws.cell(2, 2).fill = WHITE
    return ws


def _note(ws, row, text, column=2):
    cell = ws.cell(row, column, ILLEGAL_CHARACTERS_RE.sub("", text)[:32767])
    cell.data_type = "s"
    cell.font = Font(name="Malgun Gothic", size=11, color="555555")
    cell.fill = WHITE
    cell.alignment = Alignment(wrap_text=False)


def _add_image(ws, path, anchor, max_width=900):
    image = Image(path)
    ratio = min(1, max_width / image.width)
    image.width *= ratio
    image.height *= ratio
    ws.add_image(image, anchor)
    return 5 + int(image.height / 20)


def _write_100(ws, spots, factor, unit):
    sites = list(dict.fromkeys(m.site for m, values, _ in spots if values[3] is not None))
    channels = list(dict.fromkeys(m.channel for m, values, _ in spots if values[3] is not None))
    lookup = {(m.channel, m.site): values[3] * factor for m, values, _ in spots if values[3] is not None}
    rows = []
    for channel in channels:
        values = [lookup.get((channel, site)) for site in sites]
        available = [v for v in values if v is not None]
        maximum = max(available)
        average = sum(available) / len(available)
        deviation = (maximum - average) / average if average != 0 else None
        rows.append([channel, maximum, average, deviation, *values])
    _table(ws, [f"Channel / [{unit}]", "MAX.", "AVG.", "편차", *sites], rows,
           number_columns=range(1, len(sites) + 4))
    for r in range(5, len(rows) + 5):
        ws.cell(r, 5).number_format = "0.0%"
    flags = {(m.channel, m.site): interpolated[3] for m, _, interpolated in spots}
    for r, channel in enumerate(channels, 5):
        for c, site in enumerate(sites, 6):
            if flags.get((channel, site)):
                ws.cell(r, c).number_format = '0.000"*"'
    _note(ws, len(rows) + 7, "편차 = (MAX − AVG) / AVG. 빈 셀은 주파수 범위 밖이거나 변환에 실패한 값입니다.")
    _note(ws, len(rows) + 8, "* 복소 Z 선형 보간값. MAX·AVG·편차는 표시된 유효 값으로 계산합니다.")


def _write_pi(ws, spots, factor, unit, column=2):
    rows = [[m.channel, m.site, m.priority, *[v * factor if v is not None else None for v in values]]
            for m, values, _ in spots if any(v is not None for v in values)]
    _table(ws, ["Channel", "SITE", "Priority", *[f"{v} [{unit}]" for v in FREQUENCY_LABELS]],
           rows, column=column, number_columns=range(3, 8))
    valid_spots = [spot for spot in spots if any(v is not None for v in spot[1])]
    for row, (_, _, flags) in enumerate(valid_spots, 5):
        for offset, interpolated in enumerate(flags):
            if interpolated:
                ws.cell(row, column + 3 + offset).number_format = '0.000"*"'
    _note(ws, len(rows) + 7, "|Zii|. * 복소 Z 선형 보간값. 주파수 범위 밖은 공란입니다.", column)


def _write_dcr(ws, rows):
    sites = list(dict.fromkeys(row.site for row in rows))
    channels = list(dict.fromkeys(row.channel for row in rows))
    lookup = {}
    for row in rows:
        key = (row.channel, row.site)
        if key in lookup or not isfinite(row.resistance_mohm) or row.resistance_mohm < 0:
            raise ValueError("DCR 값은 채널/사이트별로 하나의 유효한 음이 아닌 값이어야 합니다.")
        lookup[key] = row.resistance_mohm
    _table(ws, ["Channel / [mΩ]", *sites],
           [[channel, *[lookup.get((channel, site)) for site in sites]] for channel in channels],
           number_columns=range(1, len(sites) + 1))


def _write_caps(ws, rows, spd, config):
    sites = list(dict.fromkeys(row["site"] for row in rows))
    channels = list(dict.fromkeys(row["channel"] for row in rows))
    values = list(dict.fromkeys(row["capacitance_f"] for row in rows))
    labels = {row["capacitance_f"]: row["label"] for row in rows}
    lookup = {(r["channel"], r["site"], r["capacitance_f"]): r["count"] for r in rows}
    mapped = {(m.channel, m.site) for m in _mappings(config)}
    row_number = 4
    for value in values:
        data = [[channel, *[lookup.get((channel, site, value), 0) if (channel, site) in mapped else None
                            for site in sites]] for channel in channels]
        _table(ws, [f"Channel / {labels[value]}", *sites], data, row=row_number)
        row_number += len(data) + 3
    nets = mapped_nets(spd, config)
    inactive = sum(c.is_cap and (not nets or bool(nets.intersection(c.nets))) and
                   not component_active(c, spd) for c in spd.components)
    _note(ws, row_number, f"활성 부품만 집계. Disabled/상태 미확인/Automatic 유효 모델 없음 제외: {inactive}개.")
    _note(ws, row_number + 1, "용량 미확인 부품은 미확인 항목으로 별도 표시하며 0F로 처리하지 않습니다.")


def build_report(config: ReportConfig, spd: SpdData | None, ts: TouchstoneData | None,
                 output_path: str | Path, progress=None, cancel=None) -> ReportResult:
    """Write a complete workbook atomically; failures never replace an existing report."""
    checkpoint(cancel)
    target = Path(output_path).expanduser().resolve()
    if target.suffix.lower() != ".xlsx":
        raise ValueError("출력 파일 확장자는 .xlsx여야 합니다.")
    if config.dcr_source_path and Path(config.dcr_source_path).resolve() == target:
        raise ValueError("DCR 입력 파일에 보고서를 덮어쓸 수 없습니다.")
    for source in (spd, ts):
        if source is not None and Path(source.path).resolve() == target:
            raise ValueError("입력 파일에 보고서를 덮어쓸 수 없습니다.")
    from .mapping import validate_mappings
    if any(m.selected for m in config.mappings):
        validate_mappings(config, ts)
    pages, omitted = _plan(config, spd, ts)
    if not pages:
        raise ValueError("보고서를 만들 유효한 데이터가 없습니다. 입력과 채널 매핑을 확인하세요.")
    factor = 1 if config.settings.board_type.upper() == "PCB" else 1000
    unit = "Ω" if factor == 1 else "mΩ"
    warnings = list(dict.fromkeys([*(spd.warnings if spd else []), *(ts.warnings if ts else [])]))
    if any(any(flags) for _, _, flags in _spot_rows(config, ts)):
        warnings.append("일부 기준 주파수 값에 복소 Z 선형 보간을 적용했습니다.")
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    temporary_output = None
    try:
        with tempfile.TemporaryDirectory(prefix="report-generator-") as images:
            for index, (kind, name, payload) in enumerate(pages):
                checkpoint(cancel)
                if progress:
                    progress(f"시트 생성: {name}", index, len(pages))
                ws = _base_sheet(wb, name, config)
                image_bottom = 4
                if kind == "100":
                    _write_100(ws, payload, factor, unit)
                elif kind == "pi":
                    _write_pi(ws, payload, factor, unit)
                elif kind == "dcr":
                    _write_dcr(ws, payload)
                elif kind == "cap":
                    _write_caps(ws, payload, spd, config)
                elif kind == "placement":
                    ws.freeze_panes = None
                    # openpyxl leaves frozen-pane selections behind; Excel 2013 rejects them.
                    ws.sheet_view.selection = [Selection(activeCell="B4", sqref="B4")]
                    for side_index, side in enumerate(payload):
                        checkpoint(cancel)
                        path = render_placement(spd, config, side, Path(images) / f"placement-{side}.png")
                        image_bottom = max(image_bottom, _add_image(ws, path, "B4" if side_index == 0 else "P4", max_width=850))
                elif kind == "group":
                    path = render_group_plot(ts, config, payload, Path(images) / f"group-{index}.png")
                    image_bottom = _add_image(ws, path, "B4")
                    _write_pi(ws, [(m, values, flags) for m, values, flags in _spot_rows(config, ts)
                                   if m.group == payload], factor, unit, column=16)
                if warnings:
                    _note(ws, max(ws.max_row + 3, image_bottom + 2, 7), "주의: " + " | ".join(warnings))
                last_column = max(ws.max_column, 30 if kind == "placement" else 23 if kind == "group" else 2)
                ws.print_area = f"B1:{get_column_letter(last_column)}{max(ws.max_row, image_bottom)}"
            checkpoint(cancel)
            handle, temporary_output = tempfile.mkstemp(prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent)
            os.close(handle)
            wb.save(temporary_output)
        checkpoint(cancel)
        os.replace(temporary_output, target)
        temporary_output = None
        if progress:
            progress("보고서 저장 완료", len(pages), len(pages))
    finally:
        wb.close()
        if temporary_output is not None:
            Path(temporary_output).unlink(missing_ok=True)
    return ReportResult(str(target), [name for _, name, _ in pages], omitted, warnings)
