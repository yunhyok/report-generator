import zipfile
import xml.etree.ElementTree as ET

import numpy as np
import openpyxl
import pytest

from report_generator.core.models import (CancelledError, ChannelMapping, Component, DcrRow,
    Port, ReportConfig, SpiceModel, SpdData, TouchstoneData)
from report_generator.core.report import FREQUENCIES, build_report, plan_pages


def data():
    ts = TouchstoneData("synthetic.s2p", {}, [Port(1, "P1"), Port(2, "P2")],
        np.array(FREQUENCIES), np.array([[1 + 1j, 2 + 2j]] * 5))
    config = ReportConfig(mappings=[
        ChannelMapping(1, site="0", channel="VDD", net="VDD_0", group="VDD"),
        ChannelMapping(2, site="1", channel="VDD", net="VDD_1", group="VDD")])
    return config, ts


def test_units_spots_formula_metrics_and_same_chart_source(tmp_path):
    config, ts = data()
    config.settings.board_type = "PCB"
    result = build_report(config, None, ts, tmp_path / "report.xlsx")
    assert result.sheets == ["PI ALL@100MHz", "PI ALL", "VDD"]
    w = openpyxl.load_workbook(result.path)
    assert w["PI ALL"]["E5"].value == pytest.approx(2**.5)
    assert w["PI ALL@100MHz"]["F5"].value == pytest.approx(2**.5)
    assert w["PI ALL@100MHz"]["C5"].value == pytest.approx(2 * 2**.5)
    assert w["PI ALL@100MHz"]["D5"].value == pytest.approx(1.5 * 2**.5)
    assert w["PI ALL@100MHz"]["E5"].value == pytest.approx(1 / 3)
    assert w["PI ALL@100MHz"]["E5"].number_format == "0.0%"
    assert w["VDD"]["S5"].value == pytest.approx(2**.5)
    assert w["VDD"]._images[0].anchor._from.col == 1
    assert w["VDD"]._images[0].anchor._from.row == 3
    assert not w["PI ALL"].sheet_view.showGridLines
    assert w["PI ALL"]["B4"].font.name == "Malgun Gothic"
    assert w["PI ALL"]["B1"].font.sz == 24
    with zipfile.ZipFile(result.path) as z:
        assert len([n for n in z.namelist() if n.startswith("xl/media/")]) == 1
        assert not any(n.startswith("xl/charts/") for n in z.namelist())
    w.close()
    config.settings.board_type = "MLO"
    result = build_report(config, None, ts, tmp_path / "mlo.xlsx")
    w = openpyxl.load_workbook(result.path, read_only=True)
    assert w["PI ALL"]["E5"].value == pytest.approx(1000 * 2**.5)
    w.close()


def test_dcr_only_and_literal_formula_strings(tmp_path):
    config = ReportConfig(dcr=[DcrRow("0", "=SUM(1,2)", 12.5)])
    result = build_report(config, None, None, tmp_path / "dcr.xlsx")
    assert result.sheets == ["DCR"]
    w = openpyxl.load_workbook(result.path)
    assert w["DCR"]["B5"].data_type == "s"
    assert w["DCR"]["B5"].value == "=SUM(1,2)"
    assert w["DCR"]["C5"].value == 12.5
    w.close()


def test_partial_and_summary_flags_and_safe_unique_names(tmp_path):
    config, ts = data()
    config.mappings[0].group = "PI ALL"
    config.mappings[1].group = "a:/?*[]" + "한" * 35
    config.mappings[1].summary = False
    pages, omitted = plan_pages(config, None, ts)
    assert pages[:3] == ["PI ALL@100MHz", "PI ALL", "PI ALL (2)"]
    assert len(pages[-1]) <= 31 and all(c not in pages[-1] for c in ":/?*[]")
    assert "DCR" in omitted
    config.mappings[0].group = config.mappings[1].group = ""
    result = build_report(config, None, ts, tmp_path / "summary.xlsx")
    assert result.sheets == ["PI ALL@100MHz", "PI ALL"]
    w = openpyxl.load_workbook(result.path)
    assert w["PI ALL"]["C5"].value == "0"
    assert w["PI ALL"]["B6"].value is None
    w.close()


def test_unavailable_frequency_never_extrapolated(tmp_path):
    config, ts = data()
    ts.frequency_hz = np.array([1e6, 1e7])
    ts.z_diagonal = ts.z_diagonal[:2]
    pages, omitted = plan_pages(config, None, ts)
    assert "PI ALL@100MHz" not in pages
    assert "PI ALL@100MHz" in omitted
    result = build_report(config, None, ts, tmp_path / "partial.xlsx")
    w = openpyxl.load_workbook(result.path)
    assert w["PI ALL"]["E5"].value is None
    assert w["PI ALL"]["F5"].value is not None
    assert w["PI ALL"]["H5"].value is None
    w.close()


def test_caps_active_override_and_placement(tmp_path):
    config = ReportConfig(mappings=[ChannelMapping(site="0", channel="VDD", net="VDD_0")])
    config.component_overrides["C1"] = 1e-6
    spd = SpdData("synthetic.spd", {}, models={"cap": SpiceModel("cap", 1e-7, valid=True)},
        components=[Component("C1", "cap", ["VDD_0", "GND"], x_mm=1, y_mm=2,
                              width_mm=1, height_mm=.5, is_cap=True),
                    Component("C2", "cap", ["VDD_0", "GND"], status="Disabled", is_cap=True),
                    Component("C3", "NULL", ["VDD_0", "GND"], status="Enabled", is_cap=True)])
    result = build_report(config, spd, None, tmp_path / "caps.xlsx")
    assert result.sheets == ["CAP_수량", "CAP_배치도면"]
    w = openpyxl.load_workbook(result.path)
    assert w["CAP_수량"]["B4"].value == "Channel / 1uF"
    assert w["CAP_수량"]["C5"].value == 1
    assert w["CAP_수량"]["B8"].value == "Channel / Unknown"
    assert w["CAP_수량"]["C9"].value == 1
    assert len(w["CAP_배치도면"]._images) == 1
    w.close()


def test_empty_and_invalid_dcr_rejected(tmp_path):
    with pytest.raises(ValueError, match="유효한 데이터"):
        build_report(ReportConfig(), None, None, tmp_path / "empty.xlsx")
    config = ReportConfig(dcr=[DcrRow("0", "VDD", float("nan"))])
    with pytest.raises(ValueError, match="DCR"):
        build_report(config, None, None, tmp_path / "invalid.xlsx")


@pytest.mark.parametrize("board_type,image_count", [("MLO", 1), ("PCB", 2)])
def test_general_component_geometry_without_mapping_is_valid_partial_report(tmp_path, board_type, image_count):
    spd = SpdData("synthetic.spd", {}, components=[
        Component("U1", "chip", x_mm=10, y_mm=20, width_mm=3, height_mm=4),
        Component("J1", "connector", x_mm=30, y_mm=20, side="Bottom")])
    config = ReportConfig()
    config.settings.board_type = board_type
    pages, omitted = plan_pages(config, spd, None)
    assert pages == ["CAP_배치도면"]
    result = build_report(config, spd, None, tmp_path / "geometry.xlsx")
    w = openpyxl.load_workbook(result.path)
    assert len(w["CAP_배치도면"]._images) == image_count
    if image_count == 2:
        assert w["CAP_배치도면"]._images[1].anchor._from.col == 15
    w.close()


def test_cancel_and_failed_save_preserve_existing_file(tmp_path, monkeypatch):
    config, ts = data()
    for mapping in config.mappings:
        mapping.group = ""
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing report")
    calls = 0
    def cancel():
        nonlocal calls
        calls += 1
        return calls >= 3
    with pytest.raises(CancelledError):
        build_report(config, None, ts, target, cancel=cancel)
    assert target.read_bytes() == b"existing report"
    def fail_save(*args, **kwargs):
        raise OSError("disk full")
    monkeypatch.setattr("report_generator.core.report.Workbook.save", fail_save)
    with pytest.raises(OSError, match="disk full"):
        build_report(config, None, ts, target)
    assert target.read_bytes() == b"existing report"
    assert list(tmp_path.iterdir()) == [target]


def test_output_cannot_replace_source(tmp_path):
    config, ts = data()
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"source")
    ts.path = str(source)
    with pytest.raises(ValueError, match="입력 파일"):
        build_report(config, None, ts, source)
    assert source.read_bytes() == b"source"


def test_unfrozen_placement_has_no_frozen_pane_selections(tmp_path):
    spd = SpdData("synthetic.spd", {}, components=[
        Component("U1", "chip", x_mm=10, y_mm=20, width_mm=3, height_mm=4)])
    result = build_report(ReportConfig(), spd, None, tmp_path / "placement.xlsx")
    with zipfile.ZipFile(result.path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    view = root.find("s:sheetViews/s:sheetView", ns)
    assert view.find("s:pane", ns) is None
    selections = view.findall("s:selection", ns)
    assert len(selections) == 1
    assert selections[0].get("pane") is None
    assert selections[0].get("activeCell") == "B4"


def test_output_cannot_replace_dcr_input(tmp_path):
    source = tmp_path / "dcr-source.xlsx"
    source.write_bytes(b"original DCR input")
    config = ReportConfig(dcr=[DcrRow("SITE0", "VDD", 1.5)], dcr_source_path=str(source))
    with pytest.raises(ValueError, match="DCR 입력 파일"):
        build_report(config, None, None, source)
    assert source.read_bytes() == b"original DCR input"


def test_interpolated_values_are_numeric_and_marked_in_every_pi_table(tmp_path):
    config, ts = data()
    ts.frequency_hz = np.array([1e3, 1e9])
    ts.z_diagonal = ts.z_diagonal[:2]
    result = build_report(config, None, ts, tmp_path / "interpolated.xlsx")
    workbook = openpyxl.load_workbook(result.path, data_only=True)
    try:
        summary = workbook["PI ALL@100MHz"]
        assert summary["F5"].number_format == '0.000"*"'
        assert isinstance(summary["F5"].value, float)
        assert summary["C5"].number_format == "0.000"
        assert workbook["PI ALL"]["E5"].number_format == '0.000"*"'
        assert workbook["PI ALL"]["H5"].number_format == '0.000"*"'
        assert workbook["PI ALL"]["I5"].number_format == "0.000"
        assert workbook["VDD"]["S5"].number_format == '0.000"*"'
        assert workbook["VDD"]["W5"].number_format == "0.000"
    finally:
        workbook.close()
